"""问题一结构与数值一致性核验脚本。"""
from __future__ import annotations
import argparse, json
from pathlib import Path
import numpy as np, pandas as pd
from openpyxl import load_workbook

def main():
 ap=argparse.ArgumentParser(); ap.add_argument("--input-dir",type=Path,default=Path(__file__).absolute().parent/"03_原始附件"); ap.add_argument("--output-dir",type=Path,default=Path(__file__).absolute().parent/"repro_runs"); a=ap.parse_args()
 root=Path(__file__).absolute().parent/"项目代码"/"A_problem1_v3_trees"; res=root/"results"; audit=root/"audit"
 score52=pd.read_csv(res/"A_problem1_v3_scores_52.csv",index_col=0); score50=pd.read_csv(res/"A_problem1_v3_scores_50.csv",index_col=0); cont=pd.read_csv(res/"A_problem1_v3_continuous_predictions.csv",index_col=0)
 assert score52.shape==(12,52), f"评分卡 CSV 形状错误: {score52.shape}"
 assert score50.shape==(12,50) and cont.shape==(12,50), "50维输出形状错误"
 assert np.issubdtype(score52.values.dtype,np.number) and np.all((score52.values>=0)&(score52.values<=6)), "分数超出0—6"
 assert np.allclose(score52.iloc[:,:50].values,score50.values), "52列前50列与50列评分卡不一致"
 assert score52.columns[-2:].tolist()==["A_acetaldehyde","F_acetaldehyde"], "乙醛外推列缺失"
 xlsx=res/"A_problem1_v3_scores.xlsx"; sheets=load_workbook(xlsx, data_only=False, read_only=True).sheetnames
 beer=[s for s in sheets if s.startswith("beer_")]; assert len(beer)==12, f"评分卡工作表数量={len(beer)}"
 mapping=pd.read_csv(root/"data"/"scorecard_mapping.csv")
 assert len(mapping)==52 and set(mapping.dimension)==set(score52.columns), "评分卡映射不是52项"
 positions=[]
 wb=load_workbook(xlsx, data_only=False, read_only=True)
 for s in beer:
  sh=wb[s]; vals=[sh[f"{r.scorecard_column}{int(r.scorecard_row)}"].value for r in mapping.itertuples()]
  positions.append(len(vals)); assert all(v is not None and 0<=int(v)<=6 for v in vals), f"{s}评分越界或缺失"
 assert all(n==52 for n in positions), f"评分位置数量={positions}"
 summary=json.loads((audit/"q1_holdout_summary.json").read_text(encoding="utf-8"))
 metrics={"MAE":summary["continuous_mae"],"R2":summary["pooled_within_dimension_r2"],"within1":summary["proxy_grade_within1_rate"],"exact":summary["proxy_grade_exact_rate"]}
 expected={"MAE":0.37082957083529,"R2":0.275857091297219,"within1":0.83,"exact":0.37833333333333335}
 checks={k:abs(metrics[k]-v)<1e-10 for k,v in expected.items()}; assert all(checks.values()), f"Q1总体指标变化: {metrics}"
 a.output_dir.mkdir(parents=True,exist_ok=True); (a.output_dir/"verify_q1.json").write_text(json.dumps({"status":"PASS","sheets":len(beer),"positions":positions,"metrics":metrics,"checks":checks},ensure_ascii=False,indent=2),encoding="utf-8")
 print(f"Q1 PASS: 12/12 sheets, 52 positions each, 50 sensory + 2 acetaldehyde; MAE={metrics['MAE']:.14f}, R2={metrics['R2']:.15f}, ±1={metrics['within1']:.2%}, exact={metrics['exact']:.4%}")
if __name__=="__main__": main()
