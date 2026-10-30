from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

CODE = Path(r"D:\MathorCup\2026年第七届MathorCup数学应用挑战赛—大数据竞赛初赛赛题\2026年MathorCup大数据挑战赛-赛道A初赛\我的作答\代码")
OUT = Path(r"D:\MathorCup\项目代码功能解构与源码汇编.docx")

GROUPS = [
    ("一 总入口 环境与结果归档", "项目入口、环境配置、统一核验和结果归档接口。", [
        "run_all.py", "verify_all.py", "verify_q1.py", "smoke_test.py", "项目代码/sync_outputs.py"
    ]),
    ("二 问题一 感官预测与评分卡", "问题一从化学特征到连续感官预测、代理等级映射和评分卡写出的完整链路。", [
        "项目代码/A_problem1_v3_trees/src/extract_inputs.py",
        "项目代码/A_problem1_v3_trees/src/benchmark_q1_models.py",
        "项目代码/A_problem1_v3_trees/src/q1_tree_predictor.py",
        "项目代码/A_problem1_v3_trees/src/solve_problem1_v3.py",
        "项目代码/A_problem1_v3_trees/src/write_scorecards.py",
        "项目代码/A_full_solution_v4/shared/q1_tree_predictor.py",
    ]),
    ("三 问题二 关联候选与稳定性选择", "问题二完成化学—感官关联、共表达网络、去冗余 Top-3 筛选和 Bootstrap 稳定性分析。", [
        "项目代码/A_full_solution_v4/q2/src/solve_q2.py",
        "项目代码/A_full_solution_v4/q2/src/verify_q2.py",
        "项目代码/A_full_solution_v4/q2/src/plot_q2_threshold_network.py",
    ]),
    ("四 问题三 风格生成与受限迁移", "问题三完成分位变换、PCA 潜空间、条件生成、派生字段重算、风格留出和迁移约束。", [
        "项目代码/A_full_solution_v4/q3/src/solve_q3_v4.py",
        "项目代码/A_full_solution_v4/q3/src/q3_c2st_global.py",
        "项目代码/A_full_solution_v4/q3/src/q3_migration_sensitivity.py",
        "项目代码/A_full_solution_v4/q3/src/verify_q3_v4.py",
        "项目代码/A_full_solution_v4/q3/src/verify_q3_sensory_sync.py",
        "项目代码/A_full_solution_v4/shared/q3_compat.py",
    ]),
    ("五 可视化与论文图表", "读取既有结果并生成问题一、问题二、问题三及综合路线图表。", [
        "可视化代码/make_competition_figures.py",
        "可视化代码/make_heatmap_q1_matplotlib_v4.py",
        "可视化代码/make_q1_dimension_audit.py",
        "可视化代码/make_q1_figures.py",
        "可视化代码/make_q1_local_contribution_beeswarm.py",
        "可视化代码/make_q2_frequency_matplotlib.py",
        "可视化代码/make_q2_shap_comparison.py",
        "可视化代码/make_q3_beer230_transfer_profile.py",
        "可视化代码/make_rest_figures.py",
    ]),
]

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', sz='4'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:'+edge; el=borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),sz); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)

def add_code(doc, rel):
    p = doc.add_paragraph()
    p.style = 'Code Block'
    run = p.add_run((CODE/rel).read_text(encoding='utf-8-sig'))
    run.font.name='Consolas'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Consolas'); run.font.size=Pt(7.5); run.font.color.rgb=RGBColor(35,35,35)

doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
styles=doc.styles
styles['Normal'].font.name='等线'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'等线'); styles['Normal'].font.size=Pt(10.5)
for name,size,color in [('Title',20,'000000'),('Heading 1',16,'000000'),('Heading 2',13,'000000'),('Heading 3',11,'000000')]:
    s=styles[name]; s.font.name='黑体'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color)
if 'Code Block' not in styles:
    cs=styles.add_style('Code Block',WD_STYLE_TYPE.PARAGRAPH)
else: cs=styles['Code Block']
cs.paragraph_format.space_before=Pt(0); cs.paragraph_format.space_after=Pt(0); cs.paragraph_format.line_spacing=1.0

p=doc.add_paragraph(); p.style='Title'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('MathorCup 赛道 A 项目功能解构与源码汇编')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('按功能板块整理的项目说明、调用关系与 Python 源码').italic=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('整理范围：代码目录及其全部子目录中的 Python 源码')

doc.add_heading('文档说明', level=1)
doc.add_paragraph('本文档将项目代码按功能而非目录位置重新组织。每个板块先说明职责、输入、输出和与其他板块的接口，再完整收录对应 Python 文件。源码保持原样，仅通过板块标题和文件路径提供导航。')
doc.add_paragraph('项目的正式结果仍写入各问题原有的 results 和 audit 目录；总归档接口项目代码/sync_outputs.py 只负责复制既有产物到项目代码/output，不参与模型计算。')

doc.add_heading('总体调用链', level=1)
doc.add_paragraph('原始附件 → run_all.py → 问题一/问题二/问题三求解器 → 各题 results 与 audit → sync_outputs.py 汇总到 output → verify_all.py 和独立核验脚本 → 可视化脚本生成论文图表。')

t=doc.add_table(rows=1, cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
for i,h in enumerate(['板块','核心入口','主要输入','主要输出']):
 c=t.rows[0].cells[i]; c.text=h; set_cell_shading(c,'D9EAF7'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
for title,desc,files in GROUPS:
 row=t.add_row().cells; row[0].text=title; row[1].text=files[0]; row[2].text='原始附件/既有结果'; row[3].text='results、audit 或图表'
 for c in row: set_cell_border(c)

doc.add_heading('源码板块', level=1)
for gi,(title,desc,files) in enumerate(GROUPS):
    doc.add_page_break(); doc.add_heading(title, level=1); doc.add_paragraph(desc)
    doc.add_heading('文件清单', level=2)
    for rel in files:
        p=doc.add_paragraph(style='List Bullet'); p.add_run(rel)
    doc.add_heading('源码', level=2)
    for rel in files:
        path=CODE/rel
        if not path.exists():
            doc.add_paragraph(f'文件不存在：{rel}', style='Intense Quote'); continue
        doc.add_heading(rel, level=3)
        doc.add_paragraph(f'源文件位置：{path}', style='Caption')
        add_code(doc, rel)

# appendix inventory

doc.add_page_break(); doc.add_heading('附录 结果接口与审计接口', level=1)
doc.add_paragraph('Q1 正式结果包括连续预测、50 项代理评分、52 项评分卡和 Excel 工作簿；Q2 正式结果包括每个感官维度的 Top-3 候选及稳定性汇总；Q3 正式结果包括 1100 条生成谱、代理感官画像、11 条迁移方案及缺陷核验。audit 目录中的文件用于说明输入、模型选择、稳定性、留出表现、派生量重算和边界条件，建议随项目长期保留。')
doc.add_paragraph('运行总入口：run_all.py；只读总核验入口：verify_all.py；结果总归档接口：项目代码/sync_outputs.py。运行总入口会在三个求解程序成功后自动调用同步器，但不会自动调用可视化脚本。')

doc.save(OUT)
print(OUT)
