from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path(r"D:\MathorCup\MathorCup赛道A环境配置与运行命令指南.docx")
doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles
for name,font,size in [('Normal','等线',10.5),('Title','黑体',20),('Heading 1','黑体',15),('Heading 2','黑体',12)]:
 s=styles[name]; s.font.name=font; s._element.rPr.rFonts.set(qn('w:eastAsia'),font); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
if 'Command Block' not in styles:
 s=styles.add_style('Command Block',WD_STYLE_TYPE.PARAGRAPH)
else:s=styles['Command Block']
s.font.name='Consolas'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Consolas'); s.font.size=Pt(8.5); s.font.color.rgb=RGBColor(35,35,35); s.paragraph_format.space_before=Pt(2); s.paragraph_format.space_after=Pt(7); s.paragraph_format.line_spacing=1.0

def code(text):
 p=doc.add_paragraph(style='Command Block'); p.add_run(text); return p

def bullets(items):
 for x in items: doc.add_paragraph(x,style='List Bullet')

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('MathorCup 赛道 A 环境配置与运行命令指南')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('适用于当前代码目录、结果核验和论文图表生成').italic=True

doc.add_heading('一 项目位置与目录约定',1)
doc.add_paragraph('项目代码目录：')
code(r'D:\MathorCup\2026年第七届MathorCup数学应用挑战赛—大数据竞赛初赛赛题\2026年MathorCup大数据挑战赛-赛道A初赛\我的作答\代码')
doc.add_paragraph('以下命令默认在 PowerShell 中执行，并假定当前工作目录已经切换到上述代码目录。')
code(r'$CODE = "D:\MathorCup\2026年第七届MathorCup数学应用挑战赛—大数据竞赛初赛赛题\2026年MathorCup大数据挑战赛-赛道A初赛\我的作答\代码"\nSet-Location $CODE')
doc.add_paragraph('主要目录：')
bullets(['03_原始附件：原始输入，只读使用；','项目代码：三个问题的源程序、正式结果和审计文件；','可视化代码：论文图表生成程序及图片目录；','项目代码/output：Q1、Q2、Q3 的统一结果归档接口。'])

doc.add_heading('二 Python 环境配置',1)
doc.add_paragraph('项目提供专用虚拟环境，优先使用该环境，避免系统 Python 与项目依赖版本不一致。')
code(r'.\activate_mathorcup.ps1\npython --version\npython -m pip check')
doc.add_paragraph('预期环境约为 Python 3.11.9，且 pip check 应显示 No broken requirements found。若 PowerShell 阻止脚本执行，可在当前窗口临时放宽策略：')
code(r'Set-ExecutionPolicy -Scope Process Bypass\n.\activate_mathorcup.ps1')
doc.add_heading('依赖安装或修复',2)
code(r'python -m pip install -r .\requirements.txt\npython -m pip check')
doc.add_paragraph('核心依赖包括 numpy、pandas、scipy、scikit-learn、openpyxl 和 matplotlib。不要在运行正式结果前随意升级依赖版本。')

doc.add_heading('三 推荐运行顺序',1)
doc.add_paragraph('如果只是检查当前论文支撑材料，先运行核验，不要先运行完整求解，以免覆盖既有正式产物。')
code(r'# 1. 环境与输入检查\npython --version\npython -m pip check\n\n# 2. 只读核验现有结果\npython .\verify_all.py `\n  --input-dir .\03_原始附件 `\n  --output-dir .\repro_runs `\n  --seed 20261025\n\n# 3. 补充问题三感官同步核验\n$env:MATHOR_INPUT_DIR = (Resolve-Path .\03_原始附件).Path\npython .\项目代码\A_full_solution_v4\q3\src\verify_q3_sensory_sync.py')

doc.add_heading('四 完整重新计算',1)
doc.add_paragraph('完整求解会重新生成 Q1、Q2、Q3 的 results 和 audit 文件。只有在确认需要重算时执行。')
code(r'python .\run_all.py `\n  --input-dir .\03_原始附件 `\n  --output-dir .\repro_runs `\n  --seed 20261025 `\n  --bootstrap 200')
doc.add_paragraph('执行顺序为：问题一 solve_problem1_v3.py → 问题二 solve_q2.py → 问题三 solve_q3_v4.py → 项目代码/sync_outputs.py。随机种子固定为 20261025。')
doc.add_paragraph('跳过问题一：')
code(r'python .\run_all.py --skip-q1 --bootstrap 200')
doc.add_paragraph('跳过问题一时，Q1 不会重新生成，但总归档接口仍会复制当前已有的 Q1 结果。')

doc.add_heading('五 单独运行各问题',1)
doc.add_heading('问题一',2)
code(r'python .\项目代码\A_problem1_v3_trees\src\solve_problem1_v3.py')
doc.add_heading('问题二',2)
code(r'python .\项目代码\A_full_solution_v4\q2\src\solve_q2.py `\n  --input-dir .\03_原始附件 `\n  --bootstrap 200')
doc.add_heading('问题三',2)
code(r'$env:MATHOR_INPUT_DIR = (Resolve-Path .\03_原始附件).Path\npython .\项目代码\A_full_solution_v4\q3\src\solve_q3_v4.py')

doc.add_heading('六 结果核验',1)
doc.add_paragraph('统一核验入口：')
code(r'python .\verify_all.py')
doc.add_paragraph('单独核验：')
code(r'python .\verify_q1.py --input-dir .\03_原始附件 --output-dir .\repro_runs\npython .\项目代码\A_full_solution_v4\q2\src\verify_q2.py\n$env:MATHOR_INPUT_DIR = (Resolve-Path .\03_原始附件).Path\npython .\项目代码\A_full_solution_v4\q3\src\verify_q3_v4.py\npython .\项目代码\A_full_solution_v4\q3\src\verify_q3_sensory_sync.py')
doc.add_paragraph('核验脚本读取既有结果。问题一核验会写入 repro_runs/verify_q1.json，问题三核验会更新 q3/audit/verification_checks.csv。')

doc.add_heading('七 统一结果归档',1)
doc.add_paragraph('不重新运行模型，只把当前已有实验产物复制到项目代码/output：')
code(r'python .\项目代码\sync_outputs.py')
doc.add_paragraph('归档结构：')
code(r'项目代码/output/\n├── Q1/results\n├── Q1/audit\n├── Q2/results\n├── Q2/audit\n├── Q3/results\n├── Q3/audit\n├── manifest.json\n└── README.md')
doc.add_paragraph('manifest.json 记录每个文件的来源、大小和 SHA-256。同步器不计算模型、不修改原 results 和 audit。')

doc.add_heading('八 可视化运行',1)
doc.add_paragraph('可视化应在正式结果已经存在后执行。批量入口：')
code(r'.\run_visualizations_manual.ps1')
doc.add_paragraph('清空旧图片后再生成：')
code(r'.\run_visualizations_manual.ps1 -Clear')
doc.add_paragraph('单独运行示例：')
code(r'python .\可视化代码\make_q1_figures.py\npython .\可视化代码\make_q1_dimension_audit.py\npython .\可视化代码\make_q2_shap_comparison.py\npython .\可视化代码\make_q3_beer230_transfer_profile.py\npython .\项目代码\A_full_solution_v4\q2\src\plot_q2_threshold_network.py --threshold 0.65')
doc.add_paragraph('图片默认写入 可视化代码/图表。部分脚本依赖 FIG_DIR 和 Q2_RESULTS 环境变量，推荐使用 run_visualizations_manual.ps1。')

doc.add_heading('九 运行结果位置',1)
bullets(['Q1 正式结果：项目代码/A_problem1_v3_trees/results；','Q1 审计：项目代码/A_problem1_v3_trees/audit；','Q2 正式结果：项目代码/A_full_solution_v4/q2/results；','Q2 审计：项目代码/A_full_solution_v4/q2/audit；','Q3 正式结果：项目代码/A_full_solution_v4/q3/results；','Q3 审计：项目代码/A_full_solution_v4/q3/audit；','统一归档：项目代码/output；','运行元数据：repro_runs/run_metadata.json。'])

doc.add_heading('十 常见问题',1)
doc.add_paragraph('1. 运行后关键结果发生变化：先查看 repro_runs/run_metadata.json，不要直接覆盖论文正式结果。')
doc.add_paragraph('2. 问题二核验缺少 q2_bootstrap_selected_indices.npz：说明问题二审计产物不完整，应从同一轮 solve 结果恢复。')
doc.add_paragraph('3. 图表脚本找不到环境变量：优先使用 run_visualizations_manual.ps1，或手动设置 FIG_DIR 和 Q2_RESULTS。')
doc.add_paragraph('4. 只想整理结果而不重算：执行 python .\\项目代码\\sync_outputs.py，不要执行 run_all.py。')
doc.add_paragraph('5. 论文附录中的 verify_outputs.py 和 make_figures.js 不是当前项目文件名，应使用 verify_all.py、run_visualizations_manual.ps1 及可视化代码目录中的 Python 脚本。')

doc.add_heading('十一 最小复现命令',1)
code(r'Set-Location $CODE\n.\activate_mathorcup.ps1\npython -m pip check\npython .\verify_all.py\npython .\项目代码\sync_outputs.py')
doc.add_paragraph('上述命令适用于只核验并整理当前结果的场景，不会重新训练模型。')

doc.save(OUT); print(OUT)
