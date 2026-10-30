from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

src=Path(r"D:\MathorCup\项目代码功能解构与源码汇编.docx")
out=src
interpretations={
'run_all.py':'统一编排三个问题的正式求解流程，固定输入目录、随机种子和 Bootstrap 次数，并在求解完成后记录运行元数据、检查关键结果差异，再调用总结果归档接口。',
'verify_all.py':'统一调用问题一、问题二和问题三的结果核验脚本。它只检查已有产物，不重新训练模型。',
'verify_q1.py':'核对问题一评分卡的工作表数量、评分位置、字段形状、数值范围和论文采用的总体留出指标。',
'smoke_test.py':'以较小计算规模重走问题二和问题三的关键步骤，用于检查依赖、输入格式和核心算法链路是否可运行。',
'项目代码/sync_outputs.py':'将三个问题已有的 results 和 audit 文件复制到项目代码/output/Q1、Q2、Q3，并生成带哈希值的 manifest；不参与模型计算。',
'项目代码/A_problem1_v3_trees/src/extract_inputs.py':'读取问题一所需的原始工作簿，抽取乙醛校准信息并生成输入结构审计文件，为后续建模提供稳定的数据接口。',
'项目代码/A_problem1_v3_trees/src/benchmark_q1_models.py':'在开发样本和交叉验证框架内比较候选树模型，输出模型选择所需的折外指标和扩展审计结果。',
'项目代码/A_problem1_v3_trees/src/q1_tree_predictor.py':'实现问题一冻结的树模型集成预测器，负责化学特征预处理、连续感官预测、分位锚定和代理等级映射。',
'项目代码/A_problem1_v3_trees/src/solve_problem1_v3.py':'问题一正式求解入口：训练并冻结树模型集成，生成 50 项连续预测、50/52 项代理评分、评分卡和留出审计。',
'项目代码/A_problem1_v3_trees/src/write_scorecards.py':'把问题一的评分结果按照模板字段和单元格映射写入 Excel 评分卡，并保留必要的格式和计算属性。',
'项目代码/A_full_solution_v4/shared/q1_tree_predictor.py':'供完整方案复用的问题一预测器，保证问题三生成谱的代理感官评分与问题一冻结模型保持同一套字段和参数。',
'项目代码/A_full_solution_v4/q2/src/solve_q2.py':'问题二正式求解入口，完成候选变量筛选、Spearman 关联、共表达网络融合、去冗余 Top-3 选择和 Bootstrap 稳定性分析。',
'项目代码/A_full_solution_v4/q2/src/verify_q2.py':'读取问题二结果和 Bootstrap 原始索引，逐项核对 150 条候选、候选资格、入选频率和数值完整性。',
'项目代码/A_full_solution_v4/q2/src/plot_q2_threshold_network.py':'根据化学相关矩阵和阈值绘制问题二网络图，并同时导出网络节点和边的审计表。',
'项目代码/A_full_solution_v4/q3/src/solve_q3_v4.py':'问题三正式求解入口，完成化学字段校准、分位变换、PCA 潜空间建模、风格条件生成、派生字段重算和跨风格迁移。',
'项目代码/A_full_solution_v4/q3/src/q3_c2st_global.py':'用分类器两样本检验比较真实化学谱和生成化学谱，提供生成分布差异的独立诊断。',
'项目代码/A_full_solution_v4/q3/src/q3_migration_sensitivity.py':'改变迁移搜索网格的步长，重复评估距离下降、迁移比例和缺陷约束，检查迁移结果对搜索分辨率的敏感性。',
'项目代码/A_full_solution_v4/q3/src/verify_q3_v4.py':'核对问题三生成谱、代理感官评分、迁移结果、缺陷约束、敏感性结果及五个派生字段的重算误差。',
'项目代码/A_full_solution_v4/q3/src/verify_q3_sensory_sync.py':'独立核对问题一冻结预测器对生成谱和迁移谱的代理评分，必要时进一步与问题一连续预测 CSV 对账。',
'项目代码/A_full_solution_v4/shared/q3_compat.py':'提供问题三使用的数学兼容函数，例如标准正态分位数的 Acklam 有理逼近，避免依赖额外运行时组件。',
'可视化代码/make_competition_figures.py':'生成跨问题汇总图，包括问题二候选频率热力图和三问统一计算路径图。',
'可视化代码/make_heatmap_q1_matplotlib_v4.py':'使用 Matplotlib 绘制问题一 12×52 评分卡热力图，强调等级分布和论文版式。',
'可视化代码/make_q1_dimension_audit.py':'绘制问题一感官维度的独立留出审计图，展示各维度误差和相关性。',
'可视化代码/make_q1_figures.py':'根据冻结的问题一结果生成评分卡展示图和留出审计图，服务于论文排版。',
'可视化代码/make_q1_local_contribution_beeswarm.py':'汇总树模型特征贡献并绘制样本级 beeswarm 图，用于解释化学变量对问题一预测的局部影响。',
'可视化代码/make_q2_frequency_matplotlib.py':'从问题二稳定性结果生成论文风格的 Top-3 Bootstrap 入选频率热力图。',
'可视化代码/make_q2_shap_comparison.py':'在同一数据源上比较 MDI 与 SHAP 特征重要性，辅助说明问题二候选筛选的解释性。',
'可视化代码/make_q3_beer230_transfer_profile.py':'绘制代表性 Beer 230 的源样本、迁移样本和目标风格画像，用于展示问题三迁移过程。',
'可视化代码/make_rest_figures.py':'批量生成论文中其余流程图、稳定性图、风格留出图、迁移图和缺陷图。',
}

def remove(p):
    el=p._element; el.getparent().remove(el); p._p=None

def insert_after(paragraph, text, style):
    new=OxmlElement('w:p'); paragraph._p.addnext(new)
    p=paragraph._parent.add_paragraph()
    # move newly-created p XML to the desired position, then remove trailing placement
    new2=p._p; new2.getparent().remove(new2); paragraph._p.addnext(new2)
    p._p=new2
    p.style=style; p.add_run(text)
    return p

doc=Document(src)
styles=doc.styles
if 'Interpretation' not in styles:
    s=styles.add_style('Interpretation',WD_STYLE_TYPE.PARAGRAPH)
else: s=styles['Interpretation']
s.font.name='等线'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'等线'); s.font.size=Pt(9.5); s.font.color.rgb=RGBColor(75,75,75)
s.paragraph_format.space_before=Pt(2); s.paragraph_format.space_after=Pt(5); s.paragraph_format.line_spacing=1.15

paras=list(doc.paragraphs)
for p in paras:
    if p.text.startswith('源文件位置：'):
        remove(p)
# refresh paragraphs and insert explanation once per source heading
for p in list(doc.paragraphs):
    if p.style.name=='Heading 3' and p.text in interpretations:
        # Avoid duplicate if rerun
        nxt=p._p.getnext()
        if nxt is not None and '功能解读：' in ''.join(n.text or '' for n in nxt.iter() if n.tag==qn('w:t')):
            continue
        insert_after(p, '功能解读：'+interpretations[p.text], 'Interpretation')
# Improve code paragraph readability
for p in doc.paragraphs:
    if p.style.name=='Code Block':
        p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.0
        for r in p.runs:
            r.font.name='Consolas'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Consolas'); r.font.size=Pt(7.5)
doc.save(out)
print(out)
