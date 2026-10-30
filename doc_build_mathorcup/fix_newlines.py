from docx import Document
from pathlib import Path
p=Path(r"D:\MathorCup\MathorCup赛道A环境配置与运行命令指南.docx")
d=Document(p)
for para in d.paragraphs:
 if para.style.name=='Command Block' and '\\n' in para.text:
  txt=para.text.replace('\\n','\n')
  for r in list(para.runs): para._p.remove(r._r)
  para.add_run(txt)
d.save(p)
print('fixed')
